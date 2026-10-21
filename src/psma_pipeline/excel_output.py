"""Workbook, CSV and Parquet output; data dictionary; workbook comparison.

Formatting rules applied to every sheet:

* bold, wrapped, coloured header row; header frozen; auto-filter;
* readable column widths;
* status cells coloured (PASS/OK green, WARNING amber, FAIL red,
  NOT RUN / NOT APPLICABLE grey);
* missing values are written as empty cells (documented in the dictionary).
"""

from __future__ import annotations

import logging
import re
from pathlib import Path
from typing import Iterable

import numpy as np
import pandas as pd

log = logging.getLogger("psma_pipeline.excel_output")

MISSING_CONVENTION = ("Empty cell = value not available or not applicable. Check the row's "
                      "'Calculation status' / 'Warning' columns (or RunIssues) for the reason.")

# (sheet regex, column regex, description, units, numbering scheme)
_DICT: list[tuple[str, str, str, str, str]] = [
    # ---- identifiers used on many sheets --------------------------------------
    (".*", r"PDB ID", "PDB entry identifier", "", ""),
    (".*", r"Model", "Model number in the mmCIF file (1 = first model)", "", ""),
    (".*", r"Model analysed", "Model number analysed for this entry", "", ""),
    (".*", r"Entity", "mmCIF entity ID", "", ""),
    (".*", r"Chain \(auth\)|Chain", "Author chain ID (auth_asym_id)", "", "author"),
    (".*", r"Chain \(label\)|Chains \(label\)", "mmCIF label_asym_id(s)", "", "label"),
    (".*", r"Chains \(auth\)", "Author chain IDs of the component", "", "author"),
    (".*", r"Calculation status", "OK / WARNING / FAIL / NOT RUN / NOT APPLICABLE for this row", "", ""),
    (".*", r"Warnings?", "Problems recorded for this row (empty = none)", "", ""),
    (".*", r"Status", "PASS / OK / WARNING / FAIL / NOT RUN / NOT APPLICABLE", "", ""),
    (".*", r"Method", "Experimental method(s) reported by RCSB", "", ""),
    (".*", r"Resolution \(Å\)", "Reported resolution", "Å", ""),
    (".*", r"Item", "Parameter or manifest key", "", ""),
    (".*", r"Value", "Value of the parameter / manifest key", "", ""),
    # ---- Summary -----------------------------------------------------------------
    ("Summary", r"Title", "Entry title", "", ""),
    ("Summary", r"Release date", "Initial release date", "YYYY-MM-DD", ""),
    ("Summary", r"R-free", "R-free reported by RCSB (X-ray only)", "fraction", ""),
    ("Summary", r"Space group", "Crystal space group (H-M symbol)", "", ""),
    ("Summary", r"Target chains \(auth\)", "Author chains of the target protein", "", "author"),
    ("Summary", r"Construct.*", "Per entity: UniProt range, non-UniProt N/C extensions, substitutions", "", "UniProt"),
    ("Summary", r"Missing residues \(UniProt numbering\)", "Unmodelled target residues per chain", "", "UniProt"),
    ("Summary", r"Missing internal residues", "Unmodelled residues between the first and last modelled residue", "", "UniProt"),
    ("Summary", r"(\w+) ASU", "Count of this ion in the asymmetric unit (deposited model)", "count", ""),
    ("Summary", r"(\w+) assembly", "Count of this ion in the chosen biological assembly (generator expansion)", "count", ""),
    ("Summary", r"Ligands of interest", "Ligands not classified as ion/glycan/additive: COMP(count in ASU)", "", ""),
    ("Summary", r"Glycans", "Glycan components and oligosaccharides: COMP(count in ASU)", "", ""),
    ("Summary", r"Glycan chains ASU", "Branched oligosaccharide chains in the ASU", "count", ""),
    ("Summary", r"Glycan chains assembly", "Branched oligosaccharide chains in the assembly", "count", ""),
    ("Summary", r"Glycan residues assembly.*", "Single sugar residues (non-branched entities) in the assembly", "count", ""),
    ("Summary", r"Additives / buffer", "Crystallisation / buffer components: COMP(count)", "", ""),
    ("Summary", r"Other polymer entities", "Non-target polymer entities (e.g. antibody chains)", "", ""),
    ("Summary", r"Assembly used", "Assembly ID analysed (and its RCSB 'details')", "", ""),
    ("Summary", r"Assembly choice reason", "Why this assembly was chosen (author → software → first)", "", ""),
    ("Summary", r"Assembly copies source", "How 'Target copies in assembly' was obtained", "", ""),
    ("Summary", r"RCSB assembly chain count", "Polymer chains in the chosen assembly according to RCSB", "count", ""),
    ("Summary", r"Polymer entities", "Number of polymer entities in the entry", "count", ""),
    ("Summary", r"RCSB oligomeric state", "RCSB global-symmetry oligomeric state of the chosen assembly", "", ""),
    ("Summary", r"Target copies in ASU", "Target chains in the asymmetric unit", "count", ""),
    ("Summary", r"Target copies in assembly", "Target copies in the chosen biological assembly", "count", ""),
    ("Summary", r"Monomer / dimer", "Oligomeric state of the target in the chosen assembly", "", ""),
    ("Summary", r"Assembly note", "Relation between ASU and assembly content", "", ""),
    ("Summary|TemplateRanking", r"Clashscore", "MolProbity clashscore (wwPDB validation; RCSB, PDBe fallback)", "clashes / 1000 atoms", ""),
    ("Summary|TemplateRanking", r"Ramachandran outliers \(%\)", "Ramachandran outliers (wwPDB validation)", "%", ""),
    ("Summary", r"Rotamer outliers \(%\)", "Side-chain rotamer outliers (wwPDB validation)", "%", ""),
    ("Summary", r"RSRZ outliers \(%\)", "Residues with RSRZ > 2 (X-ray only)", "%", ""),
    ("Summary", r"Clashscore percentile \(similar resolution\)", "PDBe percentile vs entries of similar resolution (higher = better)", "percentile", ""),
    ("Summary", r"Active-site residues modelled \(best chain\)", "Fraction of active-site residues modelled in the best chain", "fraction", ""),
    ("Summary", r"Active-site mean normalised B", "Mean within-chain B z-score of active-site residues (mutated residues excluded)", "z", ""),
    ("Summary", r"Active-site partial occupancy", "Any active-site atom with occupancy < 1", "", ""),
    ("Summary", r"Active-site residues differing from UniProt", "Active-site residues whose model amino acid differs from UniProt", "", "UniProt"),
    ("Summary", r"Glycosylated Asn \(UniProt\)", "Asn residues carrying an observed N-glycan", "", "UniProt"),
    ("Summary", r"Max ligand B ratio", "Highest ligand/environment B ratio among ligands of interest (indirect indicator, not map fit)", "ratio", ""),
    ("Summary", r"Entrance lid state", "Entry-level lid call (mixed = chains disagree)", "", ""),
    ("Summary", r"Entrance lid confidence", "Confidence of the lid call", "", ""),
    ("Summary", r"Arene-site contact.*", "A ligand of interest contacts an arene-site residue", "", "UniProt"),
    ("Summary", r"Largest ligand \(atoms\)", "Heavy atoms modelled in the largest ligand of interest", "count", ""),
    # ---- TemplateRanking ------------------------------------------------------------
    ("TemplateRanking", r"Rank", "Position within the ranking mode (1 = best)", "", ""),
    ("TemplateRanking", r"Score.*", "1 − weighted mean of scaled criterion ranks", "0–1", ""),
    ("TemplateRanking", r"Ranking mode", "Configured ranking mode", "", ""),
    ("TemplateRanking", r"component: .*", "Scaled rank for this criterion (0 = best, 1 = worst)", "0–1", ""),
    ("TemplateRanking", r"Criteria used", "Criteria available for this entry, with weights", "", ""),
    ("TemplateRanking", r"Score explanation", "Raw value → scaled component for each criterion", "", ""),
    ("TemplateRanking", r"Best UniProt coverage", "Highest UniProt coverage among target chains", "fraction", ""),
    ("TemplateRanking", r"Fewest internal missing", "Fewest unmodelled internal residues among target chains", "count", ""),
    ("TemplateRanking", r"Active-site completeness", "Best per-chain fraction of active-site residues modelled", "fraction", ""),
    # ---- Chains ------------------------------------------------------------------------
    ("Chains", r"Description", "Entity description", "", ""),
    ("Chains", r"Entity length", "Residues in the entity sequence (incl. tags)", "count", "label_seq"),
    ("Chains", r"Observed residues", "Entity residues with coordinates", "count", ""),
    ("Chains", r"UniProt range \(SIFTS\)", "UniProt segments mapped by SIFTS", "", "UniProt"),
    ("Chains", r"UniProt coverage", "Fraction of the UniProt sequence covered by the construct", "fraction", ""),
    ("Chains", r"N-terminal extension.*", "Entity residues before the first UniProt-mapped residue (tags)", "sequence", ""),
    ("Chains", r"C-terminal extension.*", "Entity residues after the last UniProt-mapped residue", "sequence", ""),
    ("Chains", r"Substitutions vs UniProt", "Point differences, <UniProt aa><number><entity aa>", "", "UniProt"),
    ("Chains", r"Mutation \(author annotation\)", "Depositor's mutation annotation", "", ""),
    ("Chains", r"Internal gaps / insertions", "Deletions, insertions or inconsistent SIFTS segments", "", "UniProt"),
    ("Chains", r"Missing residues \(UniProt numbering\)", "Unmodelled UniProt-mapped residues", "", "UniProt"),
    ("Chains", r"N missing \(UniProt\)", "Number of unmodelled UniProt-mapped residues", "count", ""),
    ("Chains", r"Missing internal \(loop\) residues", "Unmodelled residues between first and last modelled residue", "", "UniProt"),
    ("Chains", r"N missing internal", "Number of internal unmodelled residues", "count", ""),
    ("Chains", r"Missing non-UniProt \(tag\) residues", "Unmodelled residues outside UniProt segments", "count", ""),
    ("Chains", r"Missing residues \(entity label_seq\)", "All unmodelled residues", "", "label_seq"),
    ("Chains", r"Entity sequence", "Canonical one-letter entity sequence", "", "label_seq"),
    # ---- Ligands -------------------------------------------------------------------------
    ("Ligands", r"Component", "PDB chemical component ID", "", ""),
    ("Ligands", r"Name", "Component name", "", ""),
    ("Ligands|BindingSites|LigandModelFit", r"Class", "ion of interest / glycan / additive / buffer / ligand of interest", "", ""),
    ("Ligands", r"Count in ASU", "Instances in the asymmetric unit", "count", ""),
    ("Ligands", r"Count in assembly", "Instances in the chosen assembly", "count", ""),
    ("Ligands", r"Assembly", "Assembly ID used for the assembly count", "", ""),
    # ---- BindingSites ------------------------------------------------------------------
    ("BindingSites", r"Component", "PDB chemical component ID", "", ""),
    ("BindingSites", r"Residue no\.", "Author residue number of the ligand/ion", "", "author"),
    ("BindingSites", r"Atoms", "Modelled atoms in the ligand/ion", "count", ""),
    ("BindingSites", r"Cutoff \(Å\)", "Contact cut-off used", "Å", ""),
    ("BindingSites", r"Mean occupancy", "Mean atomic occupancy of the ligand/ion", "fraction", ""),
    ("BindingSites", r"Ligand mean B", "Mean B-factor of the ligand/ion", "Å²", ""),
    ("BindingSites", r"Contact-atom mean B", "Mean B-factor of contacting protein atoms", "Å²", ""),
    ("BindingSites", r"B ratio ligand/environment", "Ligand mean B / contact-atom mean B (indirect indicator only)", "ratio", ""),
    ("BindingSites", r"N protein residues", "Protein residues within the cut-off", "count", ""),
    ("BindingSites", r"Protein residues \(UniProt/chain\)", "Contacting residues: <aa><UniProt no.>/<author chain>; * = symmetry image", "", "UniProt (target) / author (other)"),
    ("BindingSites", r"Other ligands/ions in contact", "Non-protein, non-water partners", "", "author"),
    ("BindingSites", r"Waters in contact", "Unique water molecules within the cut-off (M1 fix)", "count", ""),
    ("BindingSites", r"Ligand atom–water pairs", "Ligand atom–water pairs within the cut-off (original count)", "count", ""),
    ("BindingSites", r"Symmetry-mate contact", "Contact with a crystal-lattice image", "", ""),
    ("BindingSites", r"Coordination number", "All atoms within the metal/anion cut-off, incl. carbons (original definition)", "count", ""),
    ("BindingSites", r"Donor coordination number", "N/O/S atoms (incl. water O) within the metal/anion cut-off", "count", ""),
    ("BindingSites", r"Donor atoms \(N/O/S: Å\)", "Potential ligand atoms and distances; '(alt X)' marks alternate locations", "Å", "UniProt / author"),
    ("BindingSites", r"Coordination \(atom: Å\)", "Coordinating atoms and distances", "Å", "UniProt / author"),
    ("BindingSites", r"Validation outliers in site", "Contacting residues with wwPDB validation outliers", "", "UniProt"),
    ("BindingSites", r"Residue/UniProt mismatches", "Contacting target residues whose amino acid differs from UniProt", "", "UniProt"),
    # ---- Glycosylation ---------------------------------------------------------------------
    ("Glycosylation", r"Asn \(UniProt\)", "Glycosylated Asn", "", "UniProt"),
    ("Glycosylation", r"Asn \(author\)", "Glycosylated Asn", "", "author"),
    ("Glycosylation", r"Root sugar", "Sugar linked to Asn ND2: COMP + author number", "", "author"),
    ("Glycosylation", r"Root sugar chain \(label\)", "label_asym_id of the root sugar", "", "label"),
    ("Glycosylation", r"C1-ND2 \(Å\)", "Sugar C1 – Asn ND2 distance", "Å", ""),
    ("Glycosylation", r"Root sugar mean B", "Mean B-factor of the root sugar", "Å²", ""),
    # ---- ActiveSite ----------------------------------------------------------------------------
    ("ActiveSite|ActiveSiteDefinition", r"UniProt residue", "Canonical amino acid + UniProt number", "", "UniProt"),
    ("ActiveSiteDefinition", r"UniProt number", "UniProt residue number", "", "UniProt"),
    ("ActiveSiteDefinition", r"Source", "UniProt feature or literature (config)", "", ""),
    ("ActiveSiteDefinition", r"Note", "Why a configured residue was dropped", "", ""),
    ("ActiveSite|ActiveSiteDefinition", r"Role", "Functional role(s) of the residue", "", ""),
    ("ActiveSite", r"Modelled", "Residue has coordinates in this chain", "", ""),
    ("ActiveSite", r"Residue in model", "Residue name + author number in the model", "", "author"),
    ("ActiveSite", r"Author residue number", "auth_seq_id + insertion code", "", "author"),
    ("ActiveSite", r"Label seq", "label_seq_id", "", "label_seq"),
    ("ActiveSite", r"Matches UniProt", "Model amino acid equals canonical UniProt amino acid", "", ""),
    ("ActiveSite", r"Atoms", "Modelled atoms in the residue", "count", ""),
    ("ActiveSite", r"Mean B", "Mean B-factor of the residue", "Å²", ""),
    ("ActiveSite", r"Side-chain mean B", "Mean B-factor of side-chain atoms", "Å²", ""),
    ("ActiveSite", r"Chain mean B", "Mean B-factor of the whole chain", "Å²", ""),
    ("ActiveSite", r"Normalised B \(z\)", "(residue mean B − chain mean B) / chain SD", "z", ""),
    ("ActiveSite", r"Min occupancy", "Lowest atomic occupancy in the residue", "fraction", ""),
    ("ActiveSite", r"Mean occupancy", "Mean atomic occupancy", "fraction", ""),
    ("ActiveSite", r"Alternate conformations", "Residue has alternate locations", "", ""),
    ("ActiveSite", r"Validation outliers", "wwPDB outlier types for the residue", "", ""),
    # ---- Validation --------------------------------------------------------------------------
    ("Validation", r"RCSB clashscore", "Clashscore (RCSB validation summary)", "clashes / 1000 atoms", ""),
    ("Validation", r"RCSB rama_outliers_pct", "Ramachandran outliers (RCSB)", "%", ""),
    ("Validation", r"RCSB rotamer_outliers_pct", "Rotamer outliers (RCSB)", "%", ""),
    ("Validation", r"RCSB rsrz_outliers_pct", "RSRZ outliers (RCSB; X-ray only)", "%", ""),
    ("Validation", r"RCSB dcc_rfree", "R-free recomputed by DCC (RCSB)", "fraction", ""),
    ("Validation", r"PDBe \S+ percentile \(vs similar resolution\)", "PDBe percentile vs similar resolution (higher = better)", "percentile", ""),
    ("Validation", r"PDBe \S+ percentile \(whole PDB\)", "PDBe percentile vs whole PDB (higher = better)", "percentile", ""),
    ("Validation", r"PDBe \S+", "Raw metric value from PDBe", "as metric", ""),
    ("Validation", r"Residues with outliers \(PDBe\)", "Residues with any validation outlier (model 1)", "count", ""),
    # ---- Variability -------------------------------------------------------------------------
    ("Variability", r"Reference", "Reference chain for superposition (PDB/author chain)", "", ""),
    ("Variability", r"Common CA", "CA atoms shared with the reference", "count", "UniProt"),
    ("Variability", r"Global CA RMSD \(Å\)", "CA RMSD after global superposition", "Å", ""),
    ("Variability", r"Largest deviation", "UniProt residue with the largest deviation (distance)", "Å", "UniProt"),
    ("Variability", r".+ CA RMSD \(Å\)", "Region CA RMSD after the global fit", "Å", "UniProt"),
    ("Variability", r".+ residues compared", "Residues of the region present in both chains", "count", ""),
    # ---- EntranceLid ---------------------------------------------------------------------------
    ("EntranceLid", r"Lid state", "open / closed / disordered / not in construct / ambiguous / undetermined", "", ""),
    ("EntranceLid", r"Lid confidence", "high (all votes agree) / medium / low", "", ""),
    ("EntranceLid", r"Lid votes", "Majority votes / available votes", "", ""),
    ("EntranceLid", r"Lid residues modelled", "Modelled lid residues / lid length", "", "UniProt"),
    ("EntranceLid", r"Lid helicity \(CA i,i\+3\)", "Fraction of helical CA(i)–CA(i+3) distances in the lid", "fraction", ""),
    ("EntranceLid", r"Vote: .*", "Vote of one classifier signal", "", ""),
    ("EntranceLid", r"Lid PC1 score", "Projection of superposed lid coordinates on PC1", "Å", ""),
    ("EntranceLid", r"Cluster", "Two-means cluster name on PC1", "", ""),
    ("EntranceLid", r"Calibration label \(literature\)", "Literature lid state of a calibration structure", "", ""),
    ("EntranceLid", r"Lid RMSD to .*", "Lid CA RMSD after superposition on non-lid residues", "Å", ""),
    # ---- LidVsLigand / Assemblies / CrossCheck -------------------------------------------------
    ("LidVsLigand", r"Ligand size", "Largest ligand of interest, heavy-atom bins", "", ""),
    ("LidVsLigand", r"open|closed", "Number of entries with this lid state", "count", ""),
    ("Assemblies", r"Assembly", "Assembly ID", "", ""),
    ("Assemblies", r"Used in summary", "This assembly was chosen for the Summary", "", ""),
    ("Assemblies", r"Defined by", "RCSB 'details' (author_defined_assembly, software_defined_assembly, …)", "", ""),
    ("Assemblies", r"Method details", "Assembly method (e.g. PISA)", "", ""),
    ("Assemblies", r"Oligomeric details", "RCSB oligomeric description", "", ""),
    ("Assemblies", r"Total polymer chains", "Polymer chains in the assembly", "count", ""),
    ("Assemblies", r"Target copies \(gemmi\)", "Target copies from the assembly generators", "count", ""),
    ("Assemblies", r"Target state", "Oligomeric state of the target in this assembly", "", ""),
    ("Assemblies", r"Symmetry", "Global symmetry symbol", "", ""),
    ("Assemblies", r"Stoichiometry", "Global stoichiometry", "", ""),
    ("Assemblies", r"RCSB oligomeric state", "RCSB oligomeric state", "", ""),
    ("Assemblies", r"Selection reason", "Why this assembly was chosen (chosen row only)", "", ""),
    ("CrossCheck", r"In (RCSB|PDBe|UniProt)", "Entry present in this source", "", ""),
    ("CrossCheck", r"(RCSB|PDBe|UniProt) resolution", "Resolution reported by this source", "Å", ""),
    ("CrossCheck", r"(RCSB|PDBe|UniProt) method", "Method reported by this source", "", ""),
    ("CrossCheck", r"PDBe chains \(UniProt range\)", "PDBe chain = UniProt range", "", "author / UniProt"),
    ("CrossCheck", r"UniProt chains", "UniProt cross-reference chains and ranges", "", "UniProt"),
    ("CrossCheck", r"Flags", "Discrepancies between sources ('consistent' if none)", "", ""),
    # ---- Milestone 2: interfaces ------------------------------------------------------------------
    ("Interfaces|CrystalContacts|InterfaceContacts", r"Source", "biological assembly or crystal lattice", "", ""),
    ("Interfaces|CrystalContacts", r"Assembly", "Assembly analysed (ASU if the assembly was unavailable)", "", ""),
    ("Interfaces|CrystalContacts|InterfaceContacts", r"Protomer [12]", "Protomer: <label_asym>|<operator> in the assembly / lattice", "", "label"),
    ("Interfaces|CrystalContacts|InterfaceContacts", r"Chain [12]", "Author chain of the protomer", "", "author"),
    ("Interfaces|CrystalContacts", r"Chain \(label\) [12]", "label_asym_id of the protomer", "", "label"),
    ("Interfaces|CrystalContacts", r"Copy [12]", "Assembly / lattice operator generating the protomer", "", ""),
    ("Interfaces|CrystalContacts|InterfaceContacts", r"Interface ID", "<PDB>:<protomer 1>-<protomer 2> (lattice interfaces marked 'lattice')", "", ""),
    ("Interfaces|CrystalContacts", r"Atom contacts", "Inter-protomer heavy-atom pairs within the contact cut-off", "count", ""),
    ("Interfaces|CrystalContacts", r"Residue contacts", "Inter-protomer residue pairs in contact", "count", ""),
    ("Interfaces|CrystalContacts", r"Interface residues [12] \(UniProt\)", "Interface residues of the protomer", "", "UniProt"),
    ("Interfaces|CrystalContacts", r"N interface residues [12]", "Interface residues of the protomer", "count", ""),
    ("Interfaces|CrystalContacts|Summary", r"Interface residue count", "Distinct UniProt residues in the interface (both protomers, homodimer union)", "count", "UniProt"),
    ("Interfaces|CrystalContacts", r"Non-UniProt interface residues", "Contacting residues without a SIFTS mapping (tags), kept apart", "count", ""),
    ("Interfaces|CrystalContacts", r"H-bonds \(distance-based\)", "Donor–acceptor N/O pairs within hbond_max (no angle criterion)", "count", ""),
    ("Interfaces|CrystalContacts", r"Salt bridges", "Acidic O – basic N pairs within salt_bridge_max", "count", ""),
    ("Interfaces|CrystalContacts", r"Hydrophobic contacts", "Hydrophobic side-chain C–C pairs within hydrophobic_max", "count", ""),
    ("Interfaces|CrystalContacts", r"Aromatic contacts", "Aromatic ring pairs with centroids within aromatic_centroid_max", "count", ""),
    ("Interfaces|CrystalContacts", r"BSA total \(Å²\)", "SASA(1) + SASA(2) − SASA(complex), protein heavy atoms", "Å²", ""),
    ("Interfaces|CrystalContacts", r"BSA protomer [12] \(Å²\)", "Surface of this protomer buried by the partner", "Å²", ""),
    ("Interfaces|CrystalContacts|Summary", r"Dimer interface area \(Å²\)", "Interface area = BSA total / 2", "Å²", ""),
    ("Interfaces|CrystalContacts", r"SASA protomer [12] \(Å²\)", "SASA of the isolated protomer", "Å²", ""),
    ("Interfaces|CrystalContacts", r"Fraction buried [12]", "BSA of the protomer / its SASA", "fraction", ""),
    ("Interfaces|CrystalContacts", r"(Non-)?[Pp]olar fraction of BSA", "Share of BSA from N/O (polar) or C (non-polar) atoms", "fraction", ""),
    ("Interfaces|CrystalContacts", r"SASA method", "freesasa or shrake_rupley (numpy fallback)", "", ""),
    ("Interfaces|CrystalContacts", r"Interface (mean|max) normalised B", "Within-chain B z-score of interface residues", "z", ""),
    ("Interfaces|CrystalContacts", r"Validation outliers at interface", "Interface residues with wwPDB validation outliers", "", "UniProt"),
    ("Interfaces|CrystalContacts", r"Glycans near interface", "Glycan residues within the neighbourhood distance of interface atoms", "", "author"),
    ("Interfaces|CrystalContacts", r"Ligands/ions near interface", "Other het groups within the neighbourhood distance of interface atoms", "", "author"),
    ("Interfaces|CrystalContacts", r"Interface conservation", "Variant / conservation mapping (Milestone 3)", "", ""),
    ("Interfaces|CrystalContacts|Summary", r"Canonical interface similarity", "Jaccard similarity of the interface residue set to the canonical interface", "0–1", "UniProt"),
    ("Interfaces", r"Contact-pair similarity", "Jaccard similarity of unordered UniProt contact pairs to the canonical pairs", "0–1", "UniProt"),
    ("Interfaces", r"Fraction of canonical interface recovered", "|interface ∩ canonical| / |canonical|", "fraction", ""),
    ("Interfaces|Summary", r"Noncanonical interface flag", "Interface differs from the canonical dimer (low Jaccard or small BSA)", "", ""),
    ("Interfaces", r"Noncanonical reason", "Which noncanonical criterion triggered", "", ""),
    ("Interfaces", r"Interface RMSD \(Å\)", "CA RMSD of canonical interface residues after superposing protomer 1 on the reference", "Å", ""),
    ("Interfaces", r"Protomer-2 rotation \(°\)", "Rotation needed to re-fit protomer 2 after the protomer-1 superposition", "degrees", ""),
    ("Interfaces", r"Protomer-2 translation \(Å\)", "Translation needed to re-fit protomer 2", "Å", ""),
    ("Interfaces", r"Protomer assignment", "Which protomer was matched to the reference protomer 1", "", ""),
    ("Interfaces", r"Interface residues compared", "CA atoms used for the interface RMSD", "count", ""),
    ("Interfaces|Summary", r"Interface cluster", "Hierarchical cluster of interface residue sets (1 − Jaccard, distance threshold)", "", ""),
    ("Interfaces", r"In main cluster", "Interface belongs to the largest cluster", "", ""),
    ("Interfaces", r"Main interface", "Largest-BSA target interface of the entry (used in Summary and comparison)", "", ""),
    ("Interfaces", r"Missing canonical interface residues [12]", "Canonical interface residues without coordinates in this protomer", "", "UniProt"),
    ("Interfaces|Summary", r"Interface complete", "All canonical interface residues are modelled in both protomers", "", ""),
    ("CrystalContacts", r"Lattice operator", "Space-group operation and lattice shift (e.g. 2_565)", "", ""),
    ("CrystalContacts", r"Matches canonical dimer", "Lattice interface reproduces the canonical dimer (Jaccard ≥ threshold)", "", ""),
    ("InterfaceContacts", r"res_[ab]", "Residue label (aa + UniProt number, or resname + author number '(auth)')", "", "UniProt / author"),
    ("InterfaceContacts", r"unp_[ab]", "UniProt number", "", "UniProt"),
    ("InterfaceContacts", r"resname_[ab]", "Residue name", "", ""),
    ("InterfaceContacts", r"auth_[ab]", "Author residue number + insertion code", "", "author"),
    ("InterfaceContacts", r"min_dist", "Minimum heavy-atom distance", "Å", ""),
    ("InterfaceContacts", r"atom_contacts", "Atom pairs within the contact cut-off", "count", ""),
    ("InterfaceContacts", r"hbonds|salt_bridges|hydrophobic|aromatic", "Typed interactions of this residue pair", "count", ""),
    ("InterfaceContacts", r"closest_atoms", "Atom names of the closest pair", "", ""),
    ("InterfaceComparison", r"Interface [AB]", "Interface IDs compared", "", ""),
    ("InterfaceComparison", r"Residue Jaccard|Contact-pair Jaccard", "Jaccard similarity", "0–1", "UniProt"),
    ("InterfaceComparison", r"Fraction of A recovered in B", "|A ∩ B| / |A| (residue sets)", "fraction", ""),
    ("CanonicalInterface", r"Item|Value", "Canonical interface definition and reference", "", "UniProt"),
    # ---- Milestone 2: local model fit ----------------------------------------------------------------
    ("LocalModelFit", r"Group", "Residue group (active site, lid, interface, ligands, ions, …)", "", ""),
    ("LocalModelFit", r"Method kind", "xray / em / other (selects the metric and rule)", "", ""),
    ("LocalModelFit", r"Residues in group", "Residues in the group for this entry", "count", ""),
    ("LocalModelFit", r"Residues with metric", "Residues with RSCC (X-ray) or Q-score (EM) in the report", "count", ""),
    ("LocalModelFit", r"Metric", "RSCC (X-ray) or Q-score (EM): never compared with each other", "", ""),
    ("LocalModelFit", r"(Median|Min) RSCC", "Real-space correlation coefficient", "0–1", ""),
    ("LocalModelFit", r"(Median|Min) Q-score", "Q-score (map–model resolvability)", "−1–1", ""),
    ("LocalModelFit", r"Median RSRZ", "Real-space R Z-score (X-ray)", "z", ""),
    ("LocalModelFit", r"Lowest-supported residue", "Residue with the lowest metric", "", "author"),
    ("LocalModelFit", r"Residues (strong|acceptable|weak)", "Residues per support category", "count", ""),
    ("LocalModelFit", r"Group support", "Category from the group statistic with the method-specific rule", "", ""),
    ("LocalModelFit|LigandModelFit", r"(Group|Local) support score", "Category as number (3 strong … 0 unavailable)", "0–3", ""),
    ("LigandModelFit", r"Component", "PDB chemical component ID", "", ""),
    ("LigandModelFit", r"Residue no\. \(author\)", "Author residue number of the ligand/ion", "", "author"),
    ("LigandModelFit", r"Alt\. conformation", "Alternate location code", "", ""),
    ("LigandModelFit", r"RSCC|RSRZ|Q-score", "Validation-report fit metric of the ligand/ion", "", ""),
    ("LigandModelFit", r"Occupancy \(avg\)", "Average occupancy", "fraction", ""),
    ("LigandModelFit", r"Clashes|Geometry outliers|Chirality outliers", "Counts from the validation report", "count", ""),
    ("LigandModelFit", r"B ratio ligand/environment \(indirect\)", "Indirect indicator only (not map fit)", "ratio", ""),
    ("LigandModelFit", r"Local support", "Method-specific support category", "", ""),
    ("EMMapFit", r"Schema", "wwPDB validation XML schema file", "", ""),
    ("EMMapFit", r"emdb_id|EMDB-resolution|Q_score|atom_inclusion_.*|contour_level_primary_map|.*fsc_resolution.*|.*percentile-qrelative",
     "EM map / model metric copied from the validation report (Entry attribute of the same name)", "as reported", ""),
    ("EMMapFit", r"Per-chain mean Q-score", "Mean residue Q-score per author chain", "", "author"),
    ("EMMapFit", r"Primary map|Half maps", "Availability according to the EMDB API ('unknown' if not reported)", "", ""),
    # ---- Milestone 2: summary helpers and rankings --------------------------------------------------
    ("Summary", r"Canonical dimer in crystal lattice", "Lattice operators reproducing the canonical dimer (for entries whose assembly lacks it)", "", ""),
    ("Summary", r"(Active-site|Ligand|Interface|Entrance-lid) local support", "Group support category (LocalModelFit)", "", ""),
    ("Summary", r"(Active-site|Ligand|Interface|Entrance-lid) local support score", "Group support as number for ranking (3 … 0)", "0–3", ""),
    ("Summary", r"Ligand local support \(worst ligand\)", "Lowest category among ligands of interest (LigandModelFit)", "", ""),
    ("Summary", r"Lowest-supported important residue", "Lowest-metric residue over all analysed groups", "", "author"),
    ("Summary", r"Local-fit availability", "Validation-report residues available", "", ""),
    ("Summary", r"Local-fit warning", "Groups with weak support", "", ""),
    ("Summary", r"Interface validation warning", "Validation outliers at the main interface", "", ""),
    ("Summary", r"Substitutions vs UniProt \(all chains\)", "All point differences from UniProt (engineered or natural)", "", "UniProt"),
    ("Summary", r"N substitutions vs UniProt", "Number of distinct substitutions", "count", ""),
    ("Summary", r"Has other polymer chains", "Non-target polymer entities present (e.g. antibodies)", "", ""),
    ("Summary", r"Glycosylated Asn count", "Distinct observed N-glycosylation sites", "count", ""),
    ("Summary", r"Zn site complete", "Zn in assembly = expected per protomer × protomers", "", ""),
    ("Summary", r"Validation XML schema", "wwPDB validation XML schema used (or 'unavailable')", "", ""),
    ("Summary", r"EM Q-score", "Entry-level Q-score (cryo-EM)", "−1–1", ""),
    ("TemplateRanking", r"RankingMode|Mode description", "Ranking mode and its purpose (config ranking.modes)", "", ""),
    ("TemplateRanking", r"Entries eligible", "Entries passing the mode's filters", "count", ""),
    # ---- run-level sheets ------------------------------------------------------------------------
    ("RunIssues", r"Stage", "Pipeline stage that recorded the issue", "", ""),
    ("RunIssues", r"Message", "Issue description", "", ""),
    ("RunIssues", r"Detail", "Additional key=value details", "", ""),
    ("SanityChecks", r"Check", "Sanity-check name", "", ""),
    ("SanityChecks", r"Critical", "A FAIL in this check stops the full run", "", ""),
    ("SanityChecks", r"Detail", "What was checked / found", "", ""),
    ("Software", r"Software", "Package or executable", "", ""),
    ("Software", r"Kind", "python package / executable", "", ""),
    ("Software", r"Available", "Importable / found on PATH", "", ""),
    ("Software", r"Version", "Installed version (packages only)", "", ""),
    ("Software", r"Location", "Module name or executable path", "", ""),
    ("ValidationSet", r"Category", "Validation-set category", "", ""),
    ("ValidationSet", r"Reason", "Why the entry was selected", "", ""),
    ("DataDictionary", r"Sheet|Column|Description|Units|Numbering", "Data-dictionary field", "", ""),
    ("WorkbookComparison", r"Sheet|Kind|Key|Column|Old|New|Detail", "Workbook comparison field", "", ""),
]
_COMPILED = [(re.compile(f"^(?:{s})$"), re.compile(f"^(?:{c})$"), d, u, n) for s, c, d, u, n in _DICT]


SHEET_ALIASES = {"AssessmentSummary": ["Summary"], "TemplateRankings": ["TemplateRanking", "Summary"],
                 "RankingTop": ["TemplateRanking"], "AssessmentIssues": ["RunIssues"]}


def describe_column(sheet: str, column: str) -> dict[str, str]:
    """Data-dictionary entry for a column (sheet-specific entries win over generic ones;
    derived sheets such as AssessmentSummary fall back to their parent sheet)."""
    best = None
    for name in [sheet] + SHEET_ALIASES.get(sheet, []):
        for s, c, d, u, n in _COMPILED:
            if c.match(str(column)) and s.match(name):
                specific = s.pattern != "^(?:.*)$"
                if best is None or (specific and not best[0]):
                    best = (specific, d, u, n)
        if best and best[0]:
            break
    if best is None:
        return {"Description": "UNDOCUMENTED", "Units": "", "Numbering": ""}
    return {"Description": best[1], "Units": best[2], "Numbering": best[3]}


def data_dictionary(tables: dict[str, pd.DataFrame]) -> pd.DataFrame:
    """One row per (sheet, column) with description, units and numbering scheme."""
    rows = []
    for sheet, df in tables.items():
        for col in df.columns:
            rows.append({"Sheet": sheet, "Column": str(col), **describe_column(sheet, str(col))})
    rows.append({"Sheet": "(all)", "Column": "(missing values)", "Description": MISSING_CONVENTION,
                 "Units": "", "Numbering": ""})
    return pd.DataFrame(rows)


STATUS_FILLS = {"PASS": "C6EFCE", "OK": "C6EFCE", "WARNING": "FFEB9C", "FAIL": "FFC7CE",
                "NOT RUN": "E7E6E6", "NOT APPLICABLE": "E7E6E6"}


def _clean(df: pd.DataFrame) -> pd.DataFrame:
    """Consistent missing values: NaN/None/'' -> empty cell; tuples/sets -> strings."""
    out = df.copy()
    for col in out.columns:
        if out[col].dtype == object:
            out[col] = out[col].map(lambda v: ", ".join(map(str, v)) if isinstance(v, (set, tuple, list)) else v)
    return out.replace({np.nan: None})


def write_workbook(path: Path, sheets: dict[str, pd.DataFrame], add_dictionary: bool = True) -> Path:
    """Write sheets to ``path`` with consistent formatting and a DataDictionary sheet.

    Sheet names longer than 31 characters are truncated (Excel limit) and the
    truncation is logged.
    """
    from openpyxl.formatting.rule import CellIsRule
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    sheets = dict(sheets)
    if add_dictionary:
        sheets["DataDictionary"] = data_dictionary({k: v for k, v in sheets.items() if k != "DataDictionary"})
    path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        for name, df in sheets.items():
            sname = name[:31]
            if sname != name:
                log.warning("sheet name %r truncated to %r", name, sname)
            df = _clean(df if isinstance(df, pd.DataFrame) else pd.DataFrame(df))
            df.to_excel(xw, sheet_name=sname, index=False)
            ws = xw.sheets[sname]
            ws.freeze_panes = "B2"
            if ws.max_row > 1 and ws.max_column >= 1:
                ws.auto_filter.ref = ws.dimensions
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="305496")
                cell.alignment = Alignment(wrap_text=True, vertical="top")
            for i, col in enumerate(df.columns, start=1):
                vals = [str(col)] + [str(v) for v in df[col].head(200) if v is not None]
                width = min(max(len(v) for v in vals) + 2, 60)
                ws.column_dimensions[get_column_letter(i)].width = max(width, 10)
                if re.search(r"status|^Flags$", str(col), re.I) and ws.max_row > 1:
                    rng = f"{get_column_letter(i)}2:{get_column_letter(i)}{ws.max_row}"
                    for word, colour in STATUS_FILLS.items():
                        ws.conditional_formatting.add(rng, CellIsRule(
                            operator="equal", formula=[f'"{word}"'],
                            fill=PatternFill("solid", fgColor=colour, bgColor=colour)))
    return path


def write_tables(folder: Path, sheets: dict[str, pd.DataFrame], csv: bool = True,
                 parquet: bool = True) -> list[Path]:
    """Write each table as CSV and/or Parquet; returns written paths.

    Parquet is skipped (with a log warning) if pyarrow/fastparquet is missing.
    """
    folder.mkdir(parents=True, exist_ok=True)
    written: list[Path] = []
    pq_ok = parquet
    for name, df in sheets.items():
        df = _clean(df)
        if csv:
            p = folder / f"{name}.csv"
            df.to_csv(p, index=False)
            written.append(p)
        if pq_ok:
            try:
                p = folder / f"{name}.parquet"
                df.astype({c: "string" for c in df.columns if df[c].dtype == object}).to_parquet(p, index=False)
                written.append(p)
            except ImportError:
                log.warning("Parquet output skipped: install pyarrow")
                pq_ok = False
    return written


# ---------------------------------------------------------------------------
# Workbook comparison (milestone regression check)
# ---------------------------------------------------------------------------
KEY_COLUMNS = ["PDB ID", "Model", "Chain", "Chain (auth)", "Entity", "Component", "Residue no.",
               "UniProt residue", "Assembly", "Asn (UniProt)", "Root sugar", "Item", "Ligand size"]


def _keys(df: pd.DataFrame) -> list[str]:
    return [k for k in KEY_COLUMNS if k in df.columns]


def _equal(a: object, b: object, tol: float) -> bool:
    """Numbers within ``tol``; strings equal, or equal as unordered ', '-separated lists."""
    if (a is None or (isinstance(a, float) and np.isnan(a))) and \
       (b is None or (isinstance(b, float) and np.isnan(b))):
        return True
    try:
        return abs(float(a) - float(b)) <= tol
    except (TypeError, ValueError):
        sa, sb = str(a).strip(), str(b).strip()
        return sa == sb or sorted(sa.split(", ")) == sorted(sb.split(", "))


def compare_workbooks(old: str | Path, new: str | Path, sheets: Iterable[str] | None = None,
                      tol: float = 1e-6, max_cells: int = 50) -> pd.DataFrame:
    """Differences between two workbooks, sheet by sheet.

    Reports added/removed sheets and columns, row-count changes, and value
    differences on shared columns for rows matched by key columns (PDB ID,
    chain, component, …). Use it at the end of each milestone against the
    previous workbook.

    Returns a table with columns Sheet, Kind, Key, Column, Old, New, Detail.
    """
    o, n = pd.read_excel(old, sheet_name=None), pd.read_excel(new, sheet_name=None)
    rows: list[dict] = []
    names = list(sheets) if sheets else sorted(set(o) | set(n))
    for s in names:
        if s not in n:
            rows.append({"Sheet": s, "Kind": "sheet removed"}); continue
        if s not in o:
            rows.append({"Sheet": s, "Kind": "sheet added"}); continue
        a, b = o[s], n[s]
        for c in a.columns.difference(b.columns):
            rows.append({"Sheet": s, "Kind": "column removed", "Column": c})
        for c in b.columns.difference(a.columns):
            rows.append({"Sheet": s, "Kind": "column added", "Column": c})
        if len(a) != len(b):
            rows.append({"Sheet": s, "Kind": "row count", "Old": len(a), "New": len(b)})
        keys = [k for k in _keys(a) if k in b.columns]
        if s == "Notes":
            continue
        if not len(a) and not len(b):
            rows.append({"Sheet": s, "Kind": "empty in both"}); continue
        if not keys:
            rows.append({"Sheet": s, "Kind": "not compared", "Detail": "no key columns"}); continue
        # repeated keys (e.g. several oligosaccharide rows per entry) are matched by occurrence,
        # after sorting on the shared columns so that row order does not matter
        shared = [c for c in a.columns if c in b.columns]
        a = a.loc[a[shared].astype(str).sort_values(shared).index]
        b = b.loc[b[shared].astype(str).sort_values(shared).index]
        a = a.assign(_occ=a.groupby(keys, dropna=False).cumcount())
        b = b.assign(_occ=b.groupby(keys, dropna=False).cumcount())
        a2 = a.set_index(keys + ["_occ"])
        b2 = b.set_index(keys + ["_occ"])
        common = a2.index.intersection(b2.index)
        if len(common) < len(a2):
            rows.append({"Sheet": s, "Kind": "rows only in old", "Old": len(a2) - len(common)})
        if len(common) < len(b2):
            rows.append({"Sheet": s, "Kind": "rows only in new", "New": len(b2) - len(common)})
        shown = 0
        for c in a2.columns.intersection(b2.columns):
            diff = [k for k in common if not _equal(a2.at[k, c], b2.at[k, c], tol)]
            if diff:
                rows.append({"Sheet": s, "Kind": "values differ", "Column": c,
                             "Detail": f"{len(diff)} of {len(common)} rows"})
                for k in diff[: max(0, max_cells - shown)]:
                    *kk, occ = k if isinstance(k, tuple) else (k,)
                    label = (kk[0] if len(kk) == 1 else str(tuple(kk))) + (f" #{occ + 1}" if occ else "")
                    rows.append({"Sheet": s, "Kind": "cell", "Key": label, "Column": c,
                                 "Old": a2.at[k, c], "New": b2.at[k, c]})
                    shown += 1
    return pd.DataFrame(rows, columns=["Sheet", "Kind", "Key", "Column", "Old", "New", "Detail"])
